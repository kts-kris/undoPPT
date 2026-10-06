"""test_layout.py - v3.5 regression tests: content-adaptive layout, layout lint, offline HTML, renderer coverage.

Every test here corresponds to a defect found by rendering the v3.4 demo in PowerPoint and Chrome.
"""

import copy
import json
import os
import re
import shutil
import tempfile
import unittest
import warnings

from pptx import Presentation

import cli
from core import __version__
from core.cognitive_planner import CognitivePlanner
from core.html_builder import HTML_RENDERERS, build_standalone_html
from core.layout_fit import MIN_FONT_PT, fit_title_size, pick_font, text_width_pt
from core.layout_lint import lint_pptx
from core.pptx_builder import RENDERERS, build_presentation
from core.render_check import available_renderers, check_html_layout

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SCENARIO_PROMPTS = {
    "project_charter": "智能客服业务立项答辩，申请首期预算，预期人效提升 40%",
    "annual_strategy_okr": "2027 年度战略规划与 OKR 制定，核心增长目标 30%",
    "qbr_business_review": "Q3 季度业务复盘 QBR，营收同比增长 18%，毛利率 42%，流失率 5%",
    "cross_team_alignment": "跨部门业务协同与共识拉通，产品研发运营三方依赖对齐",
    "team_headcount_review": "团队述职与 HC 编制申请，申请新增 6 个人头，人效提升 25%",
    "tech_rfc_review": "交易系统微服务技术选型 RFC 评审，对比自研与开源方案",
    "post_mortem_review": "支付网关 P0 故障复盘，宕机 47 分钟，影响 3.2 万笔订单",
    "product_launch_gtm": "新产品发布会与 GTM 上市方案，目标客群中大型企业",
    "enterprise_rfp_pitch": "大客户商务提案与竞标 RFP，面向银行核心系统升级",
    "promotion_assessment": "晋升答辩，申请高级工程师晋升，主导重构提升性能 3 倍",
    "internal_tech_talk": "内部技术分享：高并发系统缓存设计的常见反模式",
    "all_hands_rally": "全员大会战略誓师动员，下半年冲刺 10 亿营收目标",
}


def _tokens():
    with open(os.path.join(ROOT, "presets", "modern_bento.json"), encoding="utf-8") as f:
        return json.load(f)


class TestTextFitting(unittest.TestCase):

    def test_cjk_is_wider_than_latin(self):
        self.assertGreater(text_width_pt("方案", 20), text_width_pt("ab", 20))

    def test_numbers_with_percent_do_not_overestimate(self):
        # '94.8%' at 52pt must fit a 2.5in KPI card on one line (regression: '.' and '%' were over-measured)
        self.assertLess(text_width_pt("94.8%", 52, bold=True), 2.4 * 72)

    def test_long_title_shrinks_to_one_line(self):
        title = "痛点：传统被动方案与规则 RPA 难以支撑工业级自主闭环"
        size, one_line = fit_title_size(title, 11.7, 34)
        self.assertTrue(one_line)
        self.assertLess(size, 34)
        self.assertGreaterEqual(size, 22)

    def test_title_too_long_reports_wrap(self):
        _, one_line = fit_title_size("很长的标题" * 20, 11.7, 34)
        self.assertFalse(one_line)

    def test_css_font_stack_reduces_to_one_family(self):
        self.assertEqual(pick_font("PingFang SC, Inter, sans-serif"), "PingFang SC")
        self.assertEqual(pick_font(None), "PingFang SC")


class TestLayoutLint(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_demo_blueprint_is_lint_clean(self):
        out = os.path.join(self.tmp, "demo.pptx")
        build_presentation(cli.DEMO_BLUEPRINT, _tokens(), out)
        report = lint_pptx(out)
        self.assertEqual(report["findings"], [], report["findings"])
        self.assertEqual(report["slides"], len(cli.DEMO_BLUEPRINT["slides"]))

    def test_overloaded_slide_is_flagged(self):
        bad = {
            "layout_type": "bento_cards",
            "action_title": "痛点：" + "内容过载" * 30,
            "narrative_arc": "conflict",
            "cards": [
                {"tag": f"T{i}", "title": f"卡片{i}", "desc": "很长的说明文字" * 20,
                 "bullets": ["要点内容很长很长很长很长很长很长" * 3] * 9}
                for i in range(3)
            ],
        }
        out = os.path.join(self.tmp, "bad.pptx")
        build_presentation({"slides": [bad]}, _tokens(), out)
        codes = {f["code"] for f in lint_pptx(out)["findings"]}
        self.assertIn("TEXT_OVERFLOW", codes)
        self.assertIn("TITLE_WRAPS", codes)

    def test_wrapped_title_never_collides_with_cards(self):
        # v3.4 rendered a 34pt title onto a second line that ran into the cards below.
        slide = {
            "layout_type": "bento_cards",
            "action_title": "痛点：传统被动方案与规则 RPA 难以支撑工业级自主闭环",
            "subtitle": "副标题",
            "cards": [{"tag": "A", "title": "甲", "desc": "说明", "bullets": ["一", "二"]},
                      {"tag": "B", "title": "乙", "desc": "说明", "bullets": ["一", "二"]}],
        }
        out = os.path.join(self.tmp, "t.pptx")
        build_presentation({"slides": [slide]}, _tokens(), out)
        codes = {f["code"] for f in lint_pptx(out)["findings"]}
        self.assertNotIn("TEXT_CROSSES_SHAPE", codes)
        self.assertNotIn("TEXT_OVERLAP", codes)


class TestRendererCoverage(unittest.TestCase):

    def test_every_planner_layout_has_both_renderers(self):
        # v3.4 planner emitted 'kpi_dashboard' for QBR/OKR/headcount decks; neither builder knew it,
        # so those slides silently fell back to an empty bento page.
        with open(os.path.join(ROOT, "core", "cognitive_planner.py"), encoding="utf-8") as f:
            src = f.read()
        emitted = set(re.findall(r'"layout_type":\s*"([a-z_0-9]+)"', src))
        self.assertTrue(emitted)
        self.assertEqual(sorted(emitted - set(RENDERERS)), [])
        self.assertEqual(sorted(emitted - set(HTML_RENDERERS)), [])

    def test_unknown_layout_warns_instead_of_silently_falling_back(self):
        with tempfile.TemporaryDirectory() as tmp, warnings.catch_warnings(record=True) as w:
            warnings.simplefilter("always")
            build_presentation({"slides": [{"layout_type": "no_such_layout", "title": "x"}]}, _tokens(),
                               os.path.join(tmp, "u.pptx"))
        self.assertTrue(any("unknown layout_type" in str(x.message) for x in w))


class TestTwelveScenarioLayout(unittest.TestCase):
    """All 12 enterprise scenarios must build lint-clean decks with readable text."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.decks = {}
        planner = CognitivePlanner()
        for key, prompt in SCENARIO_PROMPTS.items():
            bp = planner.plan(prompt)
            out = os.path.join(cls.tmp, f"{key}.pptx")
            build_presentation(bp, _tokens(), out)
            cls.decks[key] = (bp, out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_all_decks_lint_clean(self):
        for key, (_, path) in self.decks.items():
            findings = lint_pptx(path)["findings"]
            self.assertEqual(findings, [], f"{key}: {findings}")

    def test_body_text_is_readable(self):
        """No more than 8% of characters under 11pt (v3.4 decision pages were 64%)."""
        total = small = 0
        for _, path in self.decks.values():
            prs = Presentation(path)
            for slide in prs.slides:
                for sh in slide.shapes:
                    if not sh.has_text_frame:
                        continue
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            n = len(r.text.strip())
                            size = r.font.size.pt if r.font.size else 18.0
                            total += n
                            if size < 11:
                                small += n
        self.assertGreater(total, 1000)
        self.assertLess(small / total, 0.08, f"{small}/{total} chars under 11pt")

    def test_decision_page_font_floor(self):
        bp, path = self.decks["project_charter"]
        prs = Presentation(path)
        last = prs.slides[len(prs.slides) - 1]
        sizes = [r.font.size.pt for sh in last.shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()]
        self.assertGreaterEqual(min(sizes), 10.0)
        body = [s for s in sizes if s < MIN_FONT_PT]
        self.assertLess(len(body), len(sizes) * 0.5)

    def test_cards_use_token_radius_not_default(self):
        _, path = self.decks["post_mortem_review"]
        prs = Presentation(path)
        radii = []
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.shape_type == 1 and sh.auto_shape_type == 5 and sh.height > 914400 * 1.5:
                    radii.append(sh.adjustments[0])
        self.assertTrue(radii)
        self.assertLess(max(radii), 0.12)  # python-pptx default is 0.1667 (visually a pill on tall cards)


class TestStandaloneHtml(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.path = os.path.join(cls.tmp, "p.html")
        build_standalone_html(cli.DEMO_BLUEPRINT, _tokens(), cls.path)
        with open(cls.path, encoding="utf-8") as f:
            cls.html = f.read()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_zero_external_dependencies(self):
        # The README promises "zero dependencies, double-click to present". v3.4 loaded Tailwind and fonts from CDNs.
        self.assertNotRegex(self.html, r'<script[^>]+src=["\']https?://')
        self.assertNotRegex(self.html, r'<link[^>]+href=["\']https?://')
        self.assertNotIn("@import url(", self.html)

    def test_fixed_canvas_scaled_to_viewport(self):
        self.assertIn("width:1340px;height:754px", self.html)
        self.assertIn("function fitStage", self.html)
        self.assertIn("window.addEventListener('resize', fitStage)", self.html)

    def test_slide_deep_link_and_static_mode(self):
        self.assertIn("get('slide')", self.html)
        self.assertIn("get('static') === '1'", self.html)

    def test_density_fit_present(self):
        self.assertIn("function fitBody", self.html)

    @unittest.skipUnless(available_renderers()["chrome"], "headless Chrome not available")
    def test_demo_html_has_no_overflow_at_desktop_and_phone_widths(self):
        n = len(cli.DEMO_BLUEPRINT["slides"])
        for size in ((1600, 900), (500, 900)):
            findings = check_html_layout(self.path, n, *size)
            self.assertEqual(findings, [], f"{size}: {findings}")


class TestVersion(unittest.TestCase):

    def test_version_is_3_x(self):
        self.assertRegex(__version__, r"^3\.[5-9]\.\d+$")

    def test_planner_stamps_engine_version(self):
        bp = CognitivePlanner().plan("季度业务复盘 QBR")
        self.assertEqual(bp["version"], __version__)


if __name__ == "__main__":
    unittest.main()
