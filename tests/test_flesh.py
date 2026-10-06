"""test_flesh.py - v3.7 tests: provenance, 待核 marking, ingest/cite, decision-ready demo, buzzword suggestions."""

import copy
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation

import cli
from core import provenance
from core.content_auditor import ContentAuditor
from core.html_builder import build_standalone_html
from core.ingest import extract_facts
from core.layout_lint import lint_pptx
from core.pptx_builder import build_presentation

from tests.test_layout import _tokens

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def slide(**kw):
    base = {"layout_type": "bento_cards", "action_title": "机遇：营收 1.2 亿，同比 +18%",
            "cards": [{"title": "现状", "desc": "人工渠道成本年增 22%，客户流失率 5%", "bullets": ["甲", "乙"]},
                      {"title": "目标", "desc": "人效提升 40%", "bullets": ["甲", "乙"]}]}
    base.update(kw)
    return base


class TestFigureDetection(unittest.TestCase):

    def test_claims_are_figures(self):
        figs = provenance.find_figures("营收 1.2 亿，同比 +18%，P99 < 15ms，¥8,600，缩减 90% 工时，47 分钟，30+ 家客户")
        for expected in ("1.2亿", "18%", "<15ms", "¥8,600", "90%", "47分钟", "30+"):
            self.assertIn(expected, figs)

    def test_structure_is_not_a_figure(self):
        self.assertEqual(provenance.find_figures("2026 年 Q3 的 P0 级 H1 阶段，3 个支柱，四层架构，L4 级安全"), [])

    def test_headcount_with_ge_is_a_figure(self):
        self.assertTrue(provenance.find_figures("申请 5 个人头"))


class TestScan(unittest.TestCase):

    def test_unsourced_slide(self):
        scan = provenance.scan_slide(slide())
        self.assertGreaterEqual(len(scan["unsourced"]), 4)
        self.assertEqual(len(scan["to_verify"]), len(scan["unsourced"]))

    def test_slide_level_source_covers_everything(self):
        scan = provenance.scan_slide(slide(source="notes.md:L3"))
        self.assertEqual(scan["unsourced"], [])
        self.assertEqual(scan["sources"], ["notes.md:L3"])

    def test_item_level_source_covers_only_that_item(self):
        s = slide()
        s["cards"][0]["source"] = "crm.csv:第2行"
        scan = provenance.scan_slide(s)
        texts = {f["text"] for f in scan["unsourced"]}
        self.assertNotIn("22%", texts)       # covered by the card's source
        self.assertIn("40%", texts)          # second card has none
        self.assertIn("18%", texts)          # slide title has none

    def test_todo_is_flagged_even_with_a_source(self):
        scan = provenance.scan_slide(slide(source="待补", status="todo"))
        self.assertTrue(scan["todo"])
        self.assertTrue(scan["to_verify"])

    def test_illustrative_is_not_unsourced_and_labelled(self):
        s = slide(status="illustrative")
        scan = provenance.scan_slide(s)
        self.assertEqual(scan["unsourced"], [])
        self.assertEqual(scan["to_verify"], [])
        self.assertIn("示例数据", provenance.provenance_labels(s)["footer"])

    def test_estimate_is_labelled_not_flagged(self):
        s = slide(source="销售口径", status="estimate")
        labels = provenance.provenance_labels(s)
        self.assertTrue(labels["footer"].startswith("估算"))
        self.assertEqual(labels["to_verify"], 0)

    def test_chart_data_needs_a_source(self):
        chart = {"layout_type": "data_chart", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}
        self.assertTrue(provenance.scan_slide(chart)["unsourced"])
        chart["source"] = "bi.csv"
        self.assertFalse(provenance.scan_slide(chart)["unsourced"])

    def test_invalid_status_is_ignored(self):
        self.assertIsNone(provenance.normalize_status("whatever"))


class TestAuditRules(unittest.TestCase):

    def setUp(self):
        self.auditor = ContentAuditor()

    def _audit(self, *slides):
        return self.auditor.audit({"slides": list(slides)})

    def test_unsourced_figures_warn(self):
        res = self._audit(slide())
        self.assertIn("UNSOURCED_FIGURES_P1", {f["code"] for f in res["findings"]})
        self.assertGreater(res["provenance"]["unsourced"], 0)

    def test_sourced_slide_is_clean(self):
        res = self._audit(slide(source="notes.md"))
        self.assertNotIn("UNSOURCED_FIGURES_P1", {f["code"] for f in res["findings"]})

    def test_todo_warns(self):
        res = self._audit(slide(source="x", status="todo"))
        self.assertIn("EVIDENCE_TODO_P1", {f["code"] for f in res["findings"]})

    def test_estimate_and_illustrative_are_informational(self):
        for status, code in (("estimate", "EVIDENCE_ESTIMATE_P1"), ("illustrative", "EVIDENCE_ILLUSTRATIVE_P1")):
            res = self._audit(slide(source="x", status=status))
            hit = [f for f in res["findings"] if f["code"] == code]
            self.assertTrue(hit, code)
            self.assertEqual(hit[0]["level"], "info")

    def test_deduction_is_capped(self):
        _, deduction = self.auditor._audit_provenance([slide() for _ in range(12)])
        self.assertEqual(deduction, ContentAuditor.PROVENANCE_DEDUCTION_CAP)

    def test_buzzwords_carry_a_replacement_suggestion(self):
        res = self._audit(slide(action_title="抓手：打通闭环"))
        hits = [f for f in res["findings"] if f["code"].startswith("BUZZWORD_DETECTED")]
        self.assertTrue(hits)
        self.assertTrue(hits[0]["suggestion"])
        self.assertIn("建议", hits[0]["message"])


class TestDecisionReadyDemo(unittest.TestCase):

    def test_demo_has_no_warnings(self):
        res = ContentAuditor().audit(cli.DEMO_BLUEPRINT)
        warnings_ = [f["code"] for f in res["findings"] if f["level"] == "warning"]
        self.assertEqual(warnings_, [])
        self.assertGreaterEqual(res["score"], 95)

    def test_demo_ends_with_a_decision(self):
        last = cli.DEMO_BLUEPRINT["slides"][-1]
        self.assertEqual(last["layout_type"], "summary")
        self.assertGreaterEqual(len(last["options"]), 3)
        self.assertEqual(sum(1 for o in last["options"] if o["recommended"]), 1)
        self.assertTrue(last["recommendation"])
        self.assertGreaterEqual(len(last["sign_off_items"]), 3)

    def test_demo_figures_are_labelled_as_illustrative(self):
        totals = provenance.scan_blueprint(cli.DEMO_BLUEPRINT)["totals"]
        self.assertGreater(totals["illustrative"], 30)
        self.assertEqual(totals["to_verify"], 0)


class TestRendering(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _shapes(self, path):
        prs = Presentation(path)
        return prs, {sh.name: sh for sl in prs.slides for sh in sl.shapes if sh.name.startswith("undoppt-")}

    def test_pptx_footer_badge_and_notes(self):
        bp = {"slides": [slide(), slide(source="notes.md:L3")]}
        out = os.path.join(self.tmp, "a.pptx")
        build_presentation(bp, _tokens(), out)
        prs = Presentation(out)
        first, second = prs.slides[0], prs.slides[1]
        names1 = [sh.name for sh in first.shapes]
        names2 = [sh.name for sh in second.shapes]
        self.assertIn("undoppt-badge", names1)
        self.assertNotIn("undoppt-badge", names2)
        self.assertIn("undoppt-footer", names2)
        self.assertIn("【待核", first.notes_slide.notes_text_frame.text)
        self.assertIn("【数据出处", second.notes_slide.notes_text_frame.text)
        self.assertIn("notes.md:L3", second.notes_slide.notes_text_frame.text)

    def test_badges_can_be_hidden_for_delivery(self):
        out = os.path.join(self.tmp, "b.pptx")
        build_presentation({"slides": [slide()]}, _tokens(), out, show_provenance_badges=False)
        self.assertNotIn("undoppt-badge", [sh.name for sh in Presentation(out).slides[0].shapes])

    def test_provenance_shapes_do_not_break_layout_lint(self):
        out = os.path.join(self.tmp, "c.pptx")
        build_presentation(cli.DEMO_BLUEPRINT, _tokens(), out)
        self.assertEqual(lint_pptx(out)["findings"], [])

    def test_provenance_shapes_do_not_shift_the_content_block(self):
        plain = {"slides": [slide(source="x")]}
        out1 = os.path.join(self.tmp, "d1.pptx")
        build_presentation(plain, _tokens(), out1)
        bare = copy.deepcopy(plain)
        del bare["slides"][0]["source"]
        out2 = os.path.join(self.tmp, "d2.pptx")
        build_presentation(bare, _tokens(), out2, show_provenance_badges=False)

        def tops(path):
            return sorted(sh.top for sh in Presentation(path).slides[0].shapes if not sh.name.startswith("undoppt-"))
        self.assertEqual(tops(out1), tops(out2))

    def test_html_footer_badge_and_drawer(self):
        out = os.path.join(self.tmp, "a.html")
        build_standalone_html({"slides": [slide(), slide(source="notes.md:L3")]}, _tokens(), out)
        with open(out, encoding="utf-8") as f:
            page = f.read()
        self.assertEqual(page.count('class="provenance-badge'), 1)
        self.assertEqual(page.count('class="provenance-footer'), 1)
        self.assertIn('data-sources="notes.md:L3"', page)
        self.assertIn('id="note-sources"', page)
        out2 = os.path.join(self.tmp, "b.html")
        build_standalone_html({"slides": [slide()]}, _tokens(), out2, show_provenance_badges=False)
        with open(out2, encoding="utf-8") as f:
            self.assertNotIn('class="provenance-badge', f.read())


class TestIngest(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _write(self, name, text):
        path = os.path.join(self.tmp, name)
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        return path

    def test_markdown_facts_carry_file_and_line(self):
        path = self._write("notes.md", "# 立项\n- Q3 营收 1.2 亿，同比 +18%\n- 客户流失率 5%\n")
        facts = {f["value"]: f for f in extract_facts(path)}
        self.assertEqual(facts["1.2亿"]["source"], "notes.md:L2")
        self.assertEqual(facts["5%"]["source"], "notes.md:L3")
        self.assertEqual(facts["18%"]["label"], "同比")
        self.assertEqual(facts["5%"]["section"], "立项")

    def test_csv_uses_header_units_and_row_numbers(self):
        path = self._write("fin.csv", "指标,数值(万元),占比(%)\n人工成本,1200,35\n外包费用,800,22\n")
        facts = {(f["label"], f["value"]): f for f in extract_facts(path)}
        key = ("人工成本 · 数值", "1200万元")
        self.assertIn(key, facts)
        self.assertEqual(facts[key]["source"], "fin.csv:第2行·数值(万元)")
        self.assertIn("35%", {f["value"] for f in extract_facts(path)})

    def test_excel_is_rejected_with_a_hint(self):
        path = self._write("x.xlsx", "")
        with self.assertRaises(ValueError):
            extract_facts(path)

    def test_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            extract_facts(os.path.join(self.tmp, "nope.md"))


class TestCite(unittest.TestCase):

    FACTS = [{"value": "1.2亿", "source": "notes.md:L2"}, {"value": "18%", "source": "notes.md:L2"},
             {"value": "22%", "source": "notes.md:L3"}, {"value": "5%", "source": "notes.md:L3"},
             {"value": "40%", "source": "notes.md:L5"}]

    def test_full_match_writes_slide_source(self):
        bp = {"slides": [slide()]}
        out, report = provenance.attach_sources(bp, self.FACTS)
        self.assertEqual(len(report["linked"]), 1)
        self.assertEqual(provenance.scan_slide(out["slides"][0])["unsourced"], [])
        self.assertIn("notes.md:L2", out["slides"][0]["source"])

    def test_partial_match_never_credits_unmatched_figures(self):
        bp = {"slides": [slide(action_title="机遇：营收 1.2 亿，另有无出处的 99%")]}
        out, report = provenance.attach_sources(bp, self.FACTS)
        s = out["slides"][0]
        self.assertNotIn("source", s)
        self.assertTrue(s["source_candidates"])
        self.assertEqual(len(report["partial"]), 1)
        self.assertIn("99%", {f["text"] for f in provenance.scan_slide(s)["unsourced"]})
        self.assertEqual([u["figure"] for u in report["unmatched"]], ["99%"])

    def test_cite_does_not_mutate_input_or_set_status(self):
        bp = {"slides": [slide()]}
        before = copy.deepcopy(bp)
        out, _ = provenance.attach_sources(bp, self.FACTS)
        self.assertEqual(bp, before)
        self.assertNotIn("status", out["slides"][0])


class TestCli(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _bp(self, name, bp):
        path = os.path.join(self.tmp, name)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(bp, f, ensure_ascii=False)
        return path

    def _run(self, *args):
        return subprocess.run([sys.executable, "cli.py", *args], cwd=ROOT, capture_output=True, text=True)

    def test_final_refuses_unsourced_figures(self):
        bp = self._bp("bad.json", {"slides": [slide()]})
        r = self._run("build", "--blueprint", bp, "--final", "--out", os.path.join(self.tmp, "o1"))
        self.assertEqual(r.returncode, 1)
        self.assertIn("--final", r.stdout)
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "o1", "presentation.pptx")))

    def test_final_builds_when_everything_is_sourced_and_hides_badges(self):
        bp = self._bp("ok.json", {"slides": [slide(source="notes.md:L2")]})
        out = os.path.join(self.tmp, "o2")
        r = self._run("build", "--blueprint", bp, "--final", "--out", out)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        names = [sh.name for sh in Presentation(os.path.join(out, "presentation.pptx")).slides[0].shapes]
        self.assertNotIn("undoppt-badge", names)

    def test_default_build_shows_badges_and_still_succeeds(self):
        bp = self._bp("draft.json", {"slides": [slide()]})
        out = os.path.join(self.tmp, "o3")
        r = self._run("build", "--blueprint", bp, "--out", out)
        self.assertEqual(r.returncode, 0)
        names = [sh.name for sh in Presentation(os.path.join(out, "presentation.pptx")).slides[0].shapes]
        self.assertIn("undoppt-badge", names)

    def test_ingest_then_cite_end_to_end(self):
        doc = os.path.join(self.tmp, "notes.md")
        with open(doc, "w", encoding="utf-8") as f:
            f.write("# 材料\n营收 1.2 亿，同比 +18%\n人工渠道成本年增 22%，客户流失率 5%\n人效提升 40%\n")
        facts = os.path.join(self.tmp, "facts.json")
        r = self._run("ingest", "--input-doc", doc, "--out", facts)
        self.assertEqual(r.returncode, 0, r.stderr)
        bp = self._bp("bp.json", {"slides": [slide()]})
        cited = os.path.join(self.tmp, "cited.json")
        r = self._run("cite", "--blueprint", bp, "--facts", facts, "--out", cited)
        self.assertEqual(r.returncode, 0, r.stderr)
        with open(cited, encoding="utf-8") as f:
            out = json.load(f)
        self.assertEqual(provenance.scan_slide(out["slides"][0])["unsourced"], [])
        r = self._run("build", "--blueprint", cited, "--final", "--out", os.path.join(self.tmp, "o4"))
        self.assertEqual(r.returncode, 0, r.stdout)

    def test_audit_prints_figure_summary(self):
        bp = self._bp("a.json", {"slides": [slide()]})
        r = self._run("audit", "--blueprint", bp)
        self.assertIn("Figures:", r.stdout)
        self.assertIn("unsourced", r.stdout)


if __name__ == "__main__":
    unittest.main()
