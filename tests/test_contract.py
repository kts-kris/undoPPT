"""test_contract.py - v3.6 tests: contract probe, evidence budget, blueprint compat, text fidelity, outline docs."""

import copy
import html as html_lib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation

from core.blueprint_compat import normalize_slide
from core.cognitive_planner import CognitivePlanner
from core.content_auditor import ContentAuditor
from core.contract_probe import SCENARIOS, probe
from core.html_builder import build_standalone_html
from core.pptx_builder import build_presentation

from tests.test_layout import SCENARIO_PROMPTS, _tokens

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class TestContractProbe(unittest.TestCase):

    def test_vague_requests_are_not_ready_and_block_on_audience_and_action(self):
        for prompt in ("做个 PPT", "帮我做一份关于 AI 的汇报"):
            r = probe(prompt)
            self.assertFalse(r["ready"], prompt)
            self.assertEqual(sorted(r["blocking"]), ["Q2", "Q4"], prompt)
            self.assertLessEqual(r["readiness"], 0.25, prompt)

    def test_blocking_questions_come_first(self):
        r = probe("智能客服业务立项答辩，申请首期预算，预期人效提升 40%")
        self.assertIn("Q2", r["blocking"])
        slots = {s["key"]: s for s in r["slots"]}
        self.assertEqual(r["questions"][0], slots["Q2"]["question"])

    def test_detailed_request_is_ready(self):
        r = probe("向 CTO 和财务总监做智能客服立项答辩：申请首期预算 300 万与 5 个人头，预期人效提升 40%，"
                  "目标 Q3 上线 MVP；备选方案是继续外包，不做的代价是客诉率持续上升")
        self.assertTrue(r["ready"], r["questions"])
        self.assertEqual(r["scenario"], "project_charter")

    def test_reference_document_counts_as_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = os.path.join(tmp, "notes.md")
            with open(doc, "w", encoding="utf-8") as f:
                f.write("# 立项\n受众是投资评审委员会和 CTO。痛点：人工渠道成本年增 22%，客户流失率上升。\n"
                        "申请预算 300 万、5 个人头；预期 ROI 40%；Q3 完成 MVP；备选方案是继续外包。")
            r = probe("智能客服业务立项答辩", doc_path=doc)
        self.assertTrue(r["ready"], r["questions"])

    def test_all_twelve_scenarios_are_classified_and_have_facts(self):
        for key, prompt in SCENARIO_PROMPTS.items():
            r = probe(prompt)
            self.assertEqual(r["scenario"], key, prompt)
            self.assertIn(key, SCENARIOS)
            self.assertGreaterEqual(len(SCENARIOS[key]["facts"]), 3)
            self.assertTrue(r["questions"] or r["ready"])

    def test_cli_probe_json(self):
        out = subprocess.run(
            [sys.executable, "cli.py", "probe", "--prompt", "做个 PPT", "--json"],
            cwd=ROOT, capture_output=True, text=True, check=True,
        ).stdout
        data = json.loads(out)
        self.assertFalse(data["ready"])
        self.assertIn("questions", data)


class TestEvidenceBudget(unittest.TestCase):

    def setUp(self):
        self.auditor = ContentAuditor()

    def _codes(self, slides):
        res = self.auditor.audit({"slides": slides})
        return {f["code"] for f in res["findings"]}

    def test_thin_columns_slide_is_flagged(self):
        slide = {"layout_type": "content_columns", "action_title": "文化：三条公约", "narrative_arc": "evidence",
                 "columns": [{"title": "甲"}, {"title": "乙"}, {"title": "丙"}]}
        self.assertIn("THIN_CONTENT_P1", self._codes([slide]))

    def test_rich_columns_slide_is_not_flagged(self):
        slide = {"layout_type": "content_columns", "action_title": "文化：三条公约",
                 "columns": [{"title": "结果论英雄", "points": ["拒绝口号自嗨与形式主义", "把精力用在解决实际业务问题上", "拿得出硬核战报才是硬道理"]},
                             {"title": "打破壁垒", "points": ["严禁跨部门推诿扯皮", "补台不拆台互相托底", "团队胜利是个人荣誉的基石"]},
                             {"title": "战功变现", "desc": "奖励只给结果", "points": ["战功与绩效奖金晋升强挂钩", "打破论资排辈能者上庸者下", "设立战功大奖当场兑现"]}]}
        self.assertNotIn("THIN_CONTENT_P1", self._codes([slide]))

    def test_kpi_slide_without_numbers_is_flagged(self):
        slide = {"layout_type": "metric_spotlight", "action_title": "战报：全面向好",
                 "metrics": [{"label": "完成率", "value": "很高", "desc": "表现优异的完成率情况说明"},
                             {"label": "留存", "value": "稳定", "desc": "客户留存保持稳定的说明文字"},
                             {"label": "成本", "value": "下降", "desc": "获客成本持续下降的说明文字"}]}
        self.assertIn("EVIDENCE_BUDGET_P1", self._codes([slide]))

    def test_kpi_slide_with_numbers_passes(self):
        slide = {"layout_type": "metric_spotlight", "action_title": "战报：达成率 104%",
                 "metrics": [{"label": "完成率", "value": "104%", "desc": "超额达成季度基准目标"},
                             {"label": "留存", "value": "91.5%", "desc": "老客复购与增购稳定"},
                             {"label": "成本", "value": "¥8,600", "desc": "获客成本同比下降"}]}
        self.assertNotIn("EVIDENCE_BUDGET_P1", self._codes([slide]))

    def test_cover_and_chart_are_exempt(self):
        cover = {"layout_type": "cover", "title": "封面"}
        chart = {"layout_type": "data_chart", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}
        codes = self._codes([cover, chart])
        self.assertFalse([c for c in codes if c.startswith("THIN_CONTENT")])

    def test_deduction_is_capped(self):
        thin = {"layout_type": "bento_cards", "action_title": "痛点：很薄", "cards": [{"title": "甲"}, {"title": "乙"}]}
        res = self.auditor.audit({"slides": [copy.deepcopy(thin) for _ in range(10)]})
        thin_codes = [f for f in res["findings"] if f["code"].startswith("THIN_CONTENT")]
        self.assertEqual(len(thin_codes), 10)  # all reported
        _, deduction = self.auditor._audit_evidence_budget([copy.deepcopy(thin) for _ in range(10)])
        self.assertEqual(deduction, ContentAuditor.EVIDENCE_DEDUCTION_CAP)


class TestBlueprintCompat(unittest.TestCase):

    def test_columns_points_and_tag_are_mapped(self):
        out = normalize_slide({"layout_type": "content_columns",
                               "columns": [{"title": "t", "tag": "铁律一", "points": ["a", "b"]}]})
        col = out["columns"][0]
        self.assertEqual(col["bullets"], ["a", "b"])
        self.assertEqual(col["badge"], "铁律一")

    def test_cross_mapping_documented_fields_are_mapped(self):
        out = normalize_slide({"layout_type": "cross_mapping", "mapping_rows": [
            {"layer": "业务层", "current": "割裂", "target": "统一", "action": "成立专班"}]})
        r = out["mapping_rows"][0]
        self.assertEqual(r["tier"], "业务层")
        self.assertEqual(r["source_desc"], "割裂")
        self.assertIn("统一", r["target_desc"])
        self.assertIn("成立专班", r["target_desc"])

    def test_horizons_and_ladder_and_matrix(self):
        h = normalize_slide({"layout_type": "horizons_curve", "horizons": [
            {"horizon": "H1", "name": "守正", "desc": "核心业务", "focus": "降本", "kpi": "毛利 55%"}]})["horizons"][0]
        self.assertEqual((h["id"], h["title"], h["metric"]), ("H1", "守正", "毛利 55%"))
        self.assertIn("核心业务", h["items"])
        lv = normalize_slide({"layout_type": "maturity_ladder", "levels": [
            {"step": "L1", "name": "起步", "desc": "d", "target": "T", "focus": "F", "metric": "M"}]})["levels"][0]
        self.assertEqual((lv["level"], lv["mechanism"]), ("L1", "F"))
        self.assertIn("T", lv["desc"])
        q = normalize_slide({"layout_type": "matrix_2x2", "axes": {"x": "广度", "y": "价值"},
                             "quadrants": [{"name": "甲", "desc": "说明", "tag": "策略"}]})
        self.assertEqual(q["x_axis"]["title"], "广度")
        self.assertEqual(q["quadrants"][0]["strategy"], "说明")
        self.assertIn("策略", q["quadrants"][0]["items"][0])

    def test_renderer_schema_is_left_untouched(self):
        slide = {"layout_type": "content_columns", "columns": [{"title": "t", "badge": "B", "bullets": ["x"]}]}
        self.assertEqual(normalize_slide(slide)["columns"][0], slide["columns"][0])

    def test_input_is_not_mutated(self):
        slide = {"layout_type": "content_columns", "columns": [{"title": "t", "points": ["a"]}]}
        before = copy.deepcopy(slide)
        normalize_slide(slide)
        self.assertEqual(slide, before)

    def test_unrelated_layouts_pass_through(self):
        slide = {"layout_type": "bento_cards", "cards": []}
        self.assertIs(normalize_slide(slide), slide)


META_KEYS = {"layout_type", "narrative_arc", "mission", "transition", "speaker_notes", "motion_pace",
             "transition_effect", "highlight", "recommended", "chart_type", "version", "core_evidence", "title"}
ZERO_LOSS_LAYOUTS = {"cross_mapping", "horizons_curve", "content_columns", "maturity_ladder", "matrix_2x2",
                     "architecture_stack", "process_flow", "timeline", "kpi_dashboard", "cover"}
# Fields intentionally not drawn on the slide face.
SKIP_PATHS = {".axes.x", ".axes.y"}


def _leaves(node, path=""):
    if isinstance(node, str):
        yield path, node
    elif isinstance(node, dict):
        for k, v in node.items():
            if k not in META_KEYS:
                yield from _leaves(v, path + "." + k)
    elif isinstance(node, list):
        for v in node:
            yield from _leaves(v, path + "[]")


def _norm(s):
    return re.sub(r"\s+", "", s)


class TestTextFidelity(unittest.TestCase):
    """Everything a blueprint says must reach the slide. v3.5 silently dropped up to 90% of some layouts."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        planner = CognitivePlanner()
        cls.decks = []
        for key, prompt in SCENARIO_PROMPTS.items():
            bp = planner.plan(prompt)
            p = os.path.join(cls.tmp, key + ".pptx")
            h = os.path.join(cls.tmp, key + ".html")
            build_presentation(bp, _tokens(), p)
            build_standalone_html(bp, _tokens(), h)
            cls.decks.append((key, bp, p, h))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def _missing(self, fmt):
        missing = []
        for key, bp, pptx_path, html_path in self.decks:
            if fmt == "pptx":
                prs = Presentation(pptx_path)
                texts = []
                for sl in prs.slides:
                    t = ""
                    for sh in sl.shapes:
                        if sh.has_text_frame:
                            t += sh.text_frame.text
                        if getattr(sh, "has_table", False) and sh.has_table:
                            t += "".join(c.text for r in sh.table.rows for c in r.cells)
                    texts.append(_norm(t))
            else:
                with open(html_path, encoding="utf-8") as f:
                    page = f.read()
                secs = re.split(r"<section ", page)[1:]
                texts = [_norm(html_lib.unescape(re.sub(r"<[^>]+>", "", s.split("</section>")[0]))) for s in secs]
            for i, sd in enumerate(bp["slides"]):
                if sd["layout_type"] not in ZERO_LOSS_LAYOUTS:
                    continue
                for path, val in _leaves(sd):
                    v = _norm(val)
                    if len(v) < 3 or path in SKIP_PATHS:
                        continue
                    if v[:8] not in texts[i]:
                        missing.append((key, i + 1, sd["layout_type"], path, val[:20]))
        return missing

    def test_pptx_carries_all_blueprint_text(self):
        missing = self._missing("pptx")
        self.assertEqual(missing, [], missing[:8])

    def test_html_carries_all_blueprint_text(self):
        missing = self._missing("html")
        self.assertEqual(missing, [], missing[:8])


class TestScenarioOutlinesDoc(unittest.TestCase):
    """docs/{en,zh}/scenario_outlines.md must describe the storylines the engine actually produces."""

    PATHS = {"zh": ("docs", "zh", "scenario_outlines.md"), "en": ("docs", "en", "scenario_outlines.md")}

    @classmethod
    def setUpClass(cls):
        cls.docs = {}
        for lang, parts in cls.PATHS.items():
            with open(os.path.join(ROOT, *parts), encoding="utf-8") as f:
                cls.docs[lang] = f.read()
        cls.doc = cls.docs["zh"]

    def _sections(self, lang="zh"):
        parts = re.split(r"\n(?=## S\d{2} )", self.docs[lang])[1:]
        out = {}
        for part in parts:
            m = re.match(r"## S\d{2} .*?`([a-z_]+)`", part)
            self.assertIsNotNone(m, part[:40])
            layouts = []
            for line in part.splitlines():
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if len(cells) >= 3 and cells[0].isdigit():
                    layouts.append(re.sub(r"[（(].*", "", cells[1]).strip())
            out[m.group(1)] = layouts
        return out

    def test_all_twelve_scenarios_documented_in_both_languages(self):
        for lang in ("zh", "en"):
            self.assertEqual(set(self._sections(lang)), set(SCENARIO_PROMPTS), lang)

    def test_outline_layouts_match_planner_storylines_in_both_languages(self):
        planner = CognitivePlanner()
        alias = {"kpi_dashboard": "metric_spotlight"}
        for lang in ("zh", "en"):
            for key, layouts in self._sections(lang).items():
                bp = planner.plan(SCENARIO_PROMPTS[key])
                planned = [alias.get(s["layout_type"], s["layout_type"]) for s in bp["slides"]]
                self.assertEqual([alias.get(x, x) for x in layouts], planned, f"{lang} {key}")

    def test_both_languages_have_the_same_page_structure(self):
        zh, en = self._sections("zh"), self._sections("en")
        self.assertEqual(zh, en)

    def test_every_outline_has_six_pages_with_required_evidence(self):
        for lang in ("zh", "en"):
            for key, layouts in self._sections(lang).items():
                self.assertEqual(len(layouts), 6, f"{lang} {key}")
        self.assertGreaterEqual(self.docs["zh"].count("必备血肉"), 12)
        self.assertGreaterEqual(self.docs["en"].count("Evidence required"), 12)

    def test_skill_md_points_to_probe_and_outlines(self):
        with open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8") as f:
            skill = f.read()
        self.assertIn("cli.py\" probe", skill)
        self.assertIn("docs/zh/scenario_outlines.md", skill)
        self.assertIn("信息不足，不生成", skill)
        with open(os.path.join(ROOT, ".agents", "skills", "undo-ppt", "SKILL.md"), encoding="utf-8") as f:
            self.assertEqual(f.read(), skill)


if __name__ == "__main__":
    unittest.main()
