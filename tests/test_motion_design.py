"""test_motion_design.py - v3.8 tests: narrative animations, contrast, real-theme extraction, template robustness."""

import copy
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation

import cli
from core import motion
from core.cognitive_planner import CognitivePlanner
from core.contrast import contrast_ratio, fix_color, rel_luminance
from core.design_check import check_tokens, repair_tokens
from core.html_builder import build_standalone_html
from core.layout_lint import lint_pptx
from core.pptx_builder import _bold_for, build_presentation
from core.render_check import available_renderers, check_motion
from core.theme_reader import read_theme
from core.undo_engine import extract_template_tokens

from tests.test_layout import SCENARIO_PROMPTS, _tokens

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def bento(highlight=True):
    return {"layout_type": "bento_cards", "action_title": "痛点：三条路线对比", "narrative_arc": "conflict",
            "cards": [{"tag": "A", "title": "甲方案", "desc": "说明文字" * 5, "bullets": ["一", "二"]},
                      {"tag": "B", "title": "乙方案", "desc": "说明文字" * 5, "bullets": ["一", "二"]},
                      {"tag": "C", "title": "丙方案", "desc": "说明文字" * 5, "bullets": ["一", "二"], "highlight": highlight}]}


class TestMotionResolution(unittest.TestCase):

    def test_off_by_default(self):
        self.assertIsNone(motion.resolve_motion(bento(), None))

    def test_deck_narrative_assigns_by_layout(self):
        self.assertEqual(motion.resolve_motion(bento(), "narrative"), "contrast")
        self.assertEqual(motion.resolve_motion({"layout_type": "metric_spotlight"}, "narrative"), "build")
        self.assertEqual(motion.resolve_motion({"layout_type": "timeline"}, "narrative"), "reveal")

    def test_cover_never_animates(self):
        self.assertIsNone(motion.resolve_motion({"layout_type": "cover", "motion": {"type": "reveal"}}, "narrative"))

    def test_slide_choice_wins_and_none_opts_out(self):
        self.assertEqual(motion.resolve_motion(dict(bento(), motion={"type": "reveal"}), None), "reveal")
        self.assertIsNone(motion.resolve_motion(dict(bento(), motion={"type": "none"}), "narrative"))

    def test_deck_off_forces_none(self):
        self.assertIsNone(motion.resolve_motion(dict(bento(), motion={"type": "reveal"}), "off"))

    def test_v33_aliases_still_work(self):
        self.assertEqual(motion.resolve_motion(dict(bento(), motion={"staged_reveal": True}), None), "reveal")
        self.assertEqual(motion.resolve_motion(dict(bento(), motion_pace="staged"), None), "contrast")
        self.assertIsNone(motion.resolve_motion(dict(bento(), motion_pace="instant"), "narrative"))


class TestNarrativeTimings(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _build(self, slides, **kw):
        out = os.path.join(self.tmp, "m.pptx")
        build_presentation({"slides": slides}, _tokens(), out, **kw)
        return Presentation(out)

    def test_no_animation_unless_asked(self):
        prs = self._build([bento()])
        self.assertFalse(motion.timing_summary(prs.slides[0])["present"])

    def test_contrast_shows_the_alternatives_first_then_the_recommended_one(self):
        prs = self._build([bento()], motion="narrative")
        slide = prs.slides[0]
        ts = motion.timing_summary(slide)
        self.assertEqual(ts["steps"], 2)
        by_id = {s.shape_id: s for s in slide.shapes}
        second_step_ids = []
        timing = [c for c in slide._element if c.tag.endswith("}timing")][0]
        steps = [c for c in timing.iter() if c.tag.endswith("}cTn") and c.get("fill") == "hold"
                 and any(k.tag.endswith("}cond") and k.get("delay") == "indefinite" for k in c.iter())]
        last = steps[-1]
        for tgt in last.iter():
            if tgt.tag.endswith("}spTgt"):
                second_step_ids.append(int(tgt.get("spid")))
        lefts = sorted({by_id[i].left for i in set(second_step_ids)})
        self.assertGreater(min(lefts), by_id[min(by_id)].left)  # the highlighted card is the rightmost column

    def test_timeline_node_and_card_animate_together(self):
        slide = {"layout_type": "timeline", "action_title": "路径：四步推进", "narrative_arc": "progression",
                 "steps": [{"time": f"Q{i}", "title": f"阶段{i}", "items": ["成果一", "成果二"]} for i in range(1, 5)]}
        prs = self._build([slide], motion="narrative")
        self.assertEqual(motion.timing_summary(prs.slides[0])["steps"], 4)

    def test_chart_then_takeaway(self):
        slide = {"layout_type": "data_chart", "action_title": "评测：领先", "categories": ["a", "b"],
                 "series": [{"name": "s", "values": [1, 2]}], "takeaway": "结论"}
        prs = self._build([slide], motion="narrative")
        self.assertEqual(motion.timing_summary(prs.slides[0])["steps"], 2)
        self.assertIn(b"bldGraphic", __import__("lxml.etree", fromlist=["x"]).tostring(prs.slides[0]._element))

    def test_single_group_slides_get_no_animation(self):
        slide = {"layout_type": "standard_table", "action_title": "规约", "headers": ["a", "b"], "rows": [["1", "2"]]}
        prs = self._build([slide], motion="narrative")
        self.assertFalse(motion.timing_summary(prs.slides[0])["present"])

    def test_header_footer_and_badge_never_animate(self):
        s = bento()
        s["figures"] = "94.8%"
        prs = self._build([s], motion="narrative")
        slide = prs.slides[0]
        animated = set(motion.timing_summary(slide)["spids"])
        for sh in slide.shapes:
            if sh.name.startswith("undoppt-") or sh.top < 1_600_000:
                self.assertNotIn(sh.shape_id, animated, sh.name)

    def test_demo_timing_trees_are_well_formed(self):
        prs = self._build(cli.DEMO_BLUEPRINT["slides"], motion="narrative")
        animated = 0
        for slide in prs.slides:
            self.assertEqual(motion.validate_timing(slide), [])
            animated += 1 if motion.timing_summary(slide)["present"] else 0
        self.assertGreaterEqual(animated, 5)

    def test_tree_follows_the_structure_powerpoint_recognises(self):
        prs = self._build([bento()], motion="narrative")
        xml = __import__("lxml.etree", fromlist=["x"]).tostring(prs.slides[0]._element).decode()
        self.assertIn('delay="indefinite"', xml)       # a click step waits for the click (v3.4 used delay="0")
        self.assertIn('nodeType="clickEffect"', xml)
        self.assertIn('nodeType="withEffect"', xml)
        self.assertIn('presetClass="entr"', xml)
        self.assertIn("<p:bldLst>", xml)

    def test_validator_catches_the_v34_mistakes(self):
        prs = self._build([bento()], motion="narrative")
        slide = prs.slides[0]
        timing = [c for c in slide._element if c.tag.endswith("}timing")][0]
        for cond in timing.iter():
            if cond.tag.endswith("}cond") and cond.get("delay") == "indefinite":
                cond.set("delay", "0")
        self.assertTrue(any("click" in p for p in motion.validate_timing(slide)))

    def test_cli_motion_flag(self):
        bp = os.path.join(self.tmp, "bp.json")
        with open(bp, "w", encoding="utf-8") as f:
            json.dump({"slides": [bento()]}, f, ensure_ascii=False)
        out = os.path.join(self.tmp, "o")
        r = subprocess.run([sys.executable, "cli.py", "build", "--blueprint", bp, "--motion", "narrative", "--out", out],
                           cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertTrue(motion.timing_summary(Presentation(os.path.join(out, "presentation.pptx")).slides[0])["present"])

    def test_html_decorative_motion_is_off_by_default(self):
        out = os.path.join(self.tmp, "a.html")
        build_standalone_html({"slides": [bento()]}, _tokens(), out)
        with open(out, encoding="utf-8") as f:
            page = f.read()
        self.assertIn('data-motion="off"', page)
        out2 = os.path.join(self.tmp, "b.html")
        build_standalone_html({"slides": [bento()]}, _tokens(), out2, motion="narrative")
        with open(out2, encoding="utf-8") as f:
            self.assertIn('data-motion="narrative"', f.read())

    @unittest.skipUnless(os.environ.get("UNDOPPT_TEST_POWERPOINT") == "1" and available_renderers()["powerpoint"],
                         "set UNDOPPT_TEST_POWERPOINT=1 on a Mac with PowerPoint to run")
    def test_powerpoint_recognises_every_animated_shape(self):
        out = os.path.join(self.tmp, "p.pptx")
        build_presentation(cli.DEMO_BLUEPRINT, _tokens(), out, motion="narrative")
        self.assertEqual(check_motion(out), [])


class TestTwelveDeckMotion(unittest.TestCase):

    def test_every_layout_in_twelve_scenarios_gets_a_valid_tree(self):
        planner = CognitivePlanner()
        with tempfile.TemporaryDirectory() as tmp:
            for key, prompt in SCENARIO_PROMPTS.items():
                out = os.path.join(tmp, key + ".pptx")
                build_presentation(planner.plan(prompt), _tokens(), out, motion="narrative")
                for i, slide in enumerate(Presentation(out).slides, 1):
                    self.assertEqual(motion.validate_timing(slide), [], f"{key} slide {i}")


class TestContrast(unittest.TestCase):

    def test_known_ratios(self):
        self.assertAlmostEqual(contrast_ratio("#000000", "#FFFFFF"), 21.0, places=1)
        self.assertAlmostEqual(contrast_ratio("#10B981", "#FFFFFF"), 2.54, places=1)

    def test_fix_color_reaches_threshold_and_keeps_hue_direction(self):
        fixed = fix_color("#10B981", "#FFFFFF", 4.5)
        self.assertGreaterEqual(contrast_ratio(fixed, "#FFFFFF"), 4.5)
        r, g, b = (int(fixed[i:i + 2], 16) for i in (0, 2, 4))
        self.assertGreater(g, r)
        self.assertGreater(g, b)  # still a green

    def test_fix_color_lightens_on_dark_backgrounds(self):
        fixed = fix_color("#334155", "#0B0C12", 4.5)
        self.assertGreater(rel_luminance(fixed), rel_luminance("#334155"))

    def test_passing_colours_are_untouched(self):
        self.assertEqual(fix_color("#0F172A", "#FFFFFF", 4.5), "0F172A")

    def test_demo_has_no_low_contrast_in_any_preset(self):
        with tempfile.TemporaryDirectory() as tmp:
            for path in sorted(glob.glob(os.path.join(ROOT, "presets", "*.json"))):
                with open(path, encoding="utf-8") as f:
                    tokens = json.load(f)
                out = os.path.join(tmp, os.path.basename(path) + ".pptx")
                build_presentation(cli.DEMO_BLUEPRINT, tokens, out)
                bad = [f for f in lint_pptx(out)["findings"] if f["code"] == "LOW_CONTRAST"]
                self.assertEqual(bad, [], os.path.basename(path))

    def test_lint_detects_low_contrast_in_a_hand_made_deck(self):
        from pptx.dml.color import RGBColor
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
        tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = "很浅的灰字"
        r.font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        from core.layout_lint import lint_contrast
        self.assertTrue(lint_contrast(s, 1))

    def test_white_text_on_a_light_primary_is_repaired(self):
        tokens = _tokens()
        tokens["palette"]["primary"] = "#38BDF8"  # light cyan: white text on it is 2.1:1
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "x.pptx")
            build_presentation(cli.DEMO_BLUEPRINT, tokens, out)
            self.assertEqual([f for f in lint_pptx(out)["findings"] if f["code"] == "LOW_CONTRAST"], [])


class TestDesignTokens(unittest.TestCase):

    def test_shipped_presets_pass(self):
        for path in sorted(glob.glob(os.path.join(ROOT, "presets", "*.json"))):
            with open(path, encoding="utf-8") as f:
                self.assertEqual(check_tokens(json.load(f)), [], os.path.basename(path))

    def test_low_contrast_and_small_sizes_are_reported_and_repaired(self):
        tokens = _tokens()
        tokens["palette"]["text_secondary"] = "#CBD5E1"
        tokens["typography"]["body"]["size"] = 9
        codes = {f["code"] for f in check_tokens(tokens)}
        self.assertIn("TOKEN_LOW_CONTRAST", codes)
        self.assertIn("TOKEN_SIZE_TOO_SMALL", codes)
        fixed, notes = repair_tokens(tokens)
        self.assertEqual(check_tokens(fixed), [])
        self.assertTrue(notes)
        self.assertEqual(tokens["typography"]["body"]["size"], 9)  # input untouched

    def test_brand_primary_is_never_rewritten(self):
        tokens = _tokens()
        tokens["palette"]["primary"] = "#F6A21D"
        fixed, _ = repair_tokens(tokens)
        self.assertEqual(fixed["palette"]["primary"], "#F6A21D")

    def test_heavy_fonts_are_not_faux_bolded(self):
        self.assertFalse(_bold_for("Haettenschweiler"))
        self.assertFalse(_bold_for("Arial Black"))
        self.assertTrue(_bold_for("PingFang SC"))


def make_template(path, scheme, clr_map, bg_xml, major="Arial", minor="Arial", hans=None):
    """A real .pptx whose theme and master are rewritten, standing in for a corporate template."""
    prs = Presentation()
    prs.save(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                text = data.decode("utf-8")
                for key, value in scheme.items():
                    text = re.sub(rf"(<a:{key}>).*?(</a:{key}>)", rf'\1<a:srgbClr val="{value}"/>\2', text, flags=re.S)
                text = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', rf"\g<1>{major}", text)
                text = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', rf"\g<1>{minor}", text)
                if hans:
                    text = re.sub(r'(<a:font script="Hans" typeface=")[^"]*', rf"\g<1>{hans}", text)
                data = text.encode("utf-8")
            elif item.filename == "ppt/slideMasters/slideMaster1.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<p:clrMap [^>]*/>", "<p:clrMap " + " ".join(f'{k}="{v}"' for k, v in clr_map.items()) + "/>", text)
                if "<p:bg>" in text:
                    text = re.sub(r"<p:bg>.*?</p:bg>", f"<p:bg>{bg_xml}</p:bg>", text, flags=re.S)
                else:
                    text = text.replace("<p:cSld>", f"<p:cSld><p:bg>{bg_xml}</p:bg>", 1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


DARK_MAP = {"bg1": "dk1", "tx1": "lt1", "bg2": "dk2", "tx2": "lt2", "accent1": "accent1", "accent2": "accent2",
            "accent3": "accent3", "accent4": "accent4", "accent5": "accent5", "accent6": "accent6",
            "hlink": "hlink", "folHlink": "folHlink"}
LIGHT_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2", "accent1": "accent1", "accent2": "accent2",
             "accent3": "accent3", "accent4": "accent4", "accent5": "accent5", "accent6": "accent6",
             "hlink": "hlink", "folHlink": "folHlink"}


class TestRealThemeExtraction(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.dark = os.path.join(self.tmp, "dark.pptx")
        make_template(self.dark, {"dk1": "000000", "lt1": "FFFFFF", "dk2": "0B0C12", "lt2": "ECF0FB",
                                  "accent1": "4970FF", "accent2": "712ED5", "accent3": "00C4B6"},
                      DARK_MAP, '<p:bgPr><a:solidFill><a:schemeClr val="bg2"/></a:solidFill><a:effectLst/></p:bgPr>',
                      major="Arial Nova Light", minor="Arial Nova Light")
        self.warm = os.path.join(self.tmp, "warm.pptx")
        make_template(self.warm, {"dk1": "000000", "lt1": "FFFFFF", "dk2": "4A5356", "lt2": "E8E3CE",
                                  "accent1": "F6A21D", "accent2": "9BAFB5", "accent3": "C96731"},
                      LIGHT_MAP,
                      '<p:bgPr><a:solidFill><a:schemeClr val="bg1"><a:lumMod val="95000"/></a:schemeClr></a:solidFill><a:effectLst/></p:bgPr>',
                      major="Gill Sans MT", minor="Gill Sans MT", hans="华文中宋")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_read_theme_resolves_the_master_background_through_the_colour_map(self):
        theme = read_theme(self.dark)
        self.assertEqual(theme["background"], "#0B0C12")  # bg2 -> dk2 in a dark master
        self.assertEqual(theme["colors"]["accent1"], "#4970FF")
        self.assertEqual(theme["fonts"]["major"], "Arial Nova Light")

    def test_read_theme_applies_lum_mod(self):
        theme = read_theme(self.warm)
        self.assertEqual(theme["background"], "#F2F2F2")  # white at 95% luminance
        self.assertEqual(theme["fonts"]["east_asian"], "华文中宋")

    def test_read_theme_returns_none_for_non_pptx(self):
        path = os.path.join(self.tmp, "x.pptx")
        with open(path, "w") as f:
            f.write("not a zip")
        self.assertIsNone(read_theme(path))

    def test_undo_detects_a_dark_template(self):
        # v3.7 returned theme_mode "light" and the default blue for every template.
        tokens = extract_template_tokens(self.dark)
        self.assertEqual(tokens["theme_mode"], "dark")
        self.assertEqual(tokens["palette"]["background"], "#0B0C12")
        self.assertEqual(tokens["palette"]["primary"], "#4970FF")
        self.assertEqual(tokens["typography"]["title"]["font"], "Arial Nova Light")
        self.assertEqual(tokens["theme_source"]["method"], "theme1.xml + slideMaster1.xml")

    def test_undo_keeps_the_brand_colour_and_cjk_font(self):
        tokens = extract_template_tokens(self.warm)
        self.assertEqual(tokens["theme_mode"], "light")
        self.assertEqual(tokens["palette"]["primary"], "#F6A21D")
        self.assertEqual(tokens["typography"]["font_ea"], "华文中宋")

    def test_two_templates_give_two_different_designs(self):
        a, b = extract_template_tokens(self.dark), extract_template_tokens(self.warm)
        self.assertNotEqual(a["palette"]["primary"], b["palette"]["primary"])
        self.assertNotEqual(a["palette"]["background"], b["palette"]["background"])

    def test_extracted_tokens_pass_the_design_check(self):
        for path in (self.dark, self.warm):
            self.assertEqual(check_tokens(extract_template_tokens(path)), [])

    def test_margins_are_clamped_so_the_header_stays_on_the_slide(self):
        for path in (self.dark, self.warm):
            canvas = extract_template_tokens(path)["canvas"]
            self.assertLessEqual(canvas["margin_left_inches"], 0.8)
            self.assertLessEqual(canvas["margin_top_inches"], 0.8)

    def test_decks_built_from_template_tokens_are_lint_clean(self):
        for name, path in (("dark", self.dark), ("warm", self.warm)):
            tokens = extract_template_tokens(path)
            out = os.path.join(self.tmp, name + "_deck.pptx")
            build_presentation(cli.DEMO_BLUEPRINT, tokens, out)
            self.assertEqual(lint_pptx(out)["findings"], [], name)

    def test_oversized_template_sizes_and_margins_cannot_break_the_header(self):
        tokens = _tokens()
        tokens["typography"]["title"]["size"] = 48
        tokens["typography"]["subtitle"]["size"] = 26
        tokens["canvas"]["margin_left_inches"] = 2.0
        tokens["canvas"]["margin_top_inches"] = 1.6
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "x.pptx")
            build_presentation(cli.DEMO_BLUEPRINT, tokens, out)
            self.assertEqual(lint_pptx(out)["findings"], [])


if __name__ == "__main__":
    unittest.main()
