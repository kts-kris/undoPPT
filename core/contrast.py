"""contrast.py - WCAG contrast maths and per-run backdrop detection for native PPTX slides (v3.8).

Used by the layout lint (report low contrast) and by the layout pass (repair it).
Thresholds: 4.5:1 for normal text, 3:1 for large text (>= 18pt, or >= 14pt bold).
"""

from typing import Any, Dict, Iterator, List, Optional, Tuple

from pptx.enum.shapes import MSO_SHAPE_TYPE

NORMAL_MIN = 4.5
LARGE_MIN = 3.0
# Measured against, not rounded up to, the threshold: 4.499 is a fail.
EPS = 1e-6


def rel_luminance(hex_color: str) -> float:
    h = hex_color.lstrip("#")
    chan = []
    for i in (0, 2, 4):
        c = int(h[i:i + 2], 16) / 255.0
        chan.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * chan[0] + 0.7152 * chan[1] + 0.0722 * chan[2]


def contrast_ratio(fg: str, bg: str) -> float:
    a, b = rel_luminance(fg), rel_luminance(bg)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


def required_ratio(size_pt: float, bold: bool) -> float:
    return LARGE_MIN if (size_pt >= 18 or (size_pt >= 14 and bold)) else NORMAL_MIN


def _blend(fg: str, target: str, amount: float) -> str:
    f, t = fg.lstrip("#"), target.lstrip("#")
    out = ""
    for i in (0, 2, 4):
        a, b = int(f[i:i + 2], 16), int(t[i:i + 2], 16)
        out += f"{round(a + (b - a) * amount):02X}"
    return out


def fix_color(fg: str, bg: str, need: float) -> str:
    """Nudge `fg` toward black or white, keeping its hue, until it reaches `need`:1 against `bg`."""
    fg = fg.lstrip("#").upper()
    bg = bg.lstrip("#").upper()
    if contrast_ratio(fg, bg) >= need - EPS:
        return fg
    toward = "000000" if rel_luminance(bg) > 0.18 else "FFFFFF"
    for step in range(1, 21):
        candidate = _blend(fg, toward, step / 20.0)
        if contrast_ratio(candidate, bg) >= need - EPS:
            return candidate
    other = "FFFFFF" if toward == "000000" else "000000"
    return max((toward, other), key=lambda c: contrast_ratio(c, bg))


def solid_fill(shape) -> Optional[str]:
    """'RRGGBB' of a shape's solid fill, or None (no fill, gradient, theme colour...)."""
    try:
        fill = shape.fill
        if fill.type == 1:  # MSO_FILL.SOLID
            return str(fill.fore_color.rgb)
    except Exception:
        return None
    return None


def run_color(run, default: str) -> str:
    try:
        if run.font.color is not None and run.font.color.type is not None:
            return str(run.font.color.rgb)
    except Exception:
        pass
    return default


def text_runs_with_backdrop(slide) -> Iterator[Tuple[Any, str, str, float, bool]]:
    """Yield (run, foreground, backdrop, size_pt, bold) for every text run on the slide."""
    shapes = list(slide.shapes)
    if not shapes:
        return
    slide_bg = solid_fill(shapes[0]) or "FFFFFF"
    fills = [(s, f) for s in shapes[1:] if (f := solid_fill(s))]

    def backdrop(shape) -> str:
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            own = solid_fill(shape)
            if own:
                return own
        cx, cy = shape.left + shape.width / 2, shape.top + shape.height / 2
        best = None
        for other, f in fills:
            if other is shape:
                continue
            if other.left <= cx <= other.left + other.width and other.top <= cy <= other.top + other.height:
                area = other.width * other.height
                if best is None or area < best[0]:
                    best = (area, f)
        return best[1] if best else slide_bg

    for shape in shapes[1:]:
        if not (shape.has_text_frame and shape.text_frame.text.strip()):
            continue
        bg = backdrop(shape)
        default_fg = "FFFFFF" if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else "000000"
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    size = r.font.size.pt if r.font.size else 18.0
                    yield r, run_color(r, default_fg), bg, size, bool(r.font.bold)
