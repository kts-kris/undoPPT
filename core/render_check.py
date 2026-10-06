"""render_check.py - Real-render verification for PPTX and HTML deliverables (v3.5).

The structural auditor scores a blueprint; it cannot see the rendered result.
This module closes that loop:

  PPTX  -> PowerPoint (macOS, via AppleScript) or LibreOffice -> PDF -> PNG -> pixel metrics
  HTML  -> headless Chrome at desktop and phone widths      -> PNG -> pixel metrics

Renderers are optional. When none is available, `RenderUnavailable` is raised and the
caller falls back to the static lint in core/layout_lint.py.
"""

import json
import os
import platform
import re
import shutil
import subprocess
import tempfile
from typing import Any, Dict, List, Optional

CHROME_PATHS = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
]
POWERPOINT_APP = "/Applications/Microsoft PowerPoint.app"
# PowerPoint is sandboxed: it can only write to folders the user has granted. ~/Documents works.
STAGING_DIR = os.path.expanduser("~/Documents/.undoppt_render")

BLANK_BAND_LIMIT = 0.35  # longest empty horizontal band, as a fraction of slide height
CONTENT_TOP_RATIO = 2.0 / 7.5


class RenderUnavailable(RuntimeError):
    pass


def available_renderers() -> Dict[str, bool]:
    return {
        "powerpoint": platform.system() == "Darwin" and os.path.isdir(POWERPOINT_APP),
        "libreoffice": bool(shutil.which("soffice") or shutil.which("libreoffice")),
        "chrome": any(os.path.exists(p) for p in CHROME_PATHS),
    }


def _require_pdfium():
    try:
        import pypdfium2  # noqa: F401
        return pypdfium2
    except ImportError as exc:
        raise RenderUnavailable("pypdfium2 is required for PDF->PNG (pip install pypdfium2 pillow)") from exc


def _pdf_to_pngs(pdf_path: str, out_dir: str, scale: float = 1.2) -> List[str]:
    pdfium = _require_pdfium()
    pdf = pdfium.PdfDocument(pdf_path)
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for i in range(len(pdf)):
        p = os.path.join(out_dir, f"slide_{i + 1:02d}.png")
        pdf[i].render(scale=scale).to_pil().save(p)
        paths.append(p)
    return paths


def _pptx_to_pdf_powerpoint(pptx_path: str, pdf_path: str, timeout: int = 150):
    os.makedirs(STAGING_DIR, exist_ok=True)
    staged = os.path.join(STAGING_DIR, "in.pptx")
    staged_pdf = os.path.join(STAGING_DIR, "out.pdf")
    shutil.copyfile(pptx_path, staged)
    if os.path.exists(staged_pdf):
        os.remove(staged_pdf)
    script = f'''
tell application "Microsoft PowerPoint"
  open POSIX file "{staged}"
  delay 2
  set p to active presentation
  save p in POSIX file "{staged_pdf}" as save as PDF
  close p saving no
end tell
'''
    try:
        subprocess.run(["osascript", "-e", script], check=True, capture_output=True, timeout=timeout)
    except subprocess.SubprocessError as exc:
        raise RenderUnavailable(f"PowerPoint export failed: {exc}") from exc
    if not os.path.exists(staged_pdf):
        raise RenderUnavailable("PowerPoint did not produce a PDF (folder access not granted?)")
    shutil.copyfile(staged_pdf, pdf_path)
    for f in (staged, staged_pdf):
        if os.path.exists(f):
            os.remove(f)


def _pptx_to_pdf_libreoffice(pptx_path: str, pdf_path: str, timeout: int = 150):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx_path],
                       check=True, capture_output=True, timeout=timeout)
        produced = os.path.join(tmp, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
        shutil.copyfile(produced, pdf_path)


def render_pptx(pptx_path: str, out_dir: str, engine: Optional[str] = None) -> Dict[str, Any]:
    """Render every slide to PNG. engine: 'powerpoint' | 'libreoffice' | None (auto)."""
    avail = available_renderers()
    order = [engine] if engine else ["powerpoint", "libreoffice"]
    os.makedirs(out_dir, exist_ok=True)
    pdf_path = os.path.join(out_dir, "deck.pdf")
    last_err = None
    for eng in order:
        if not avail.get(eng):
            continue
        try:
            if eng == "powerpoint":
                _pptx_to_pdf_powerpoint(pptx_path, pdf_path)
            else:
                _pptx_to_pdf_libreoffice(pptx_path, pdf_path)
            return {"engine": eng, "pdf": pdf_path, "pngs": _pdf_to_pngs(pdf_path, out_dir)}
        except (RenderUnavailable, subprocess.SubprocessError) as exc:
            last_err = exc
    raise RenderUnavailable(str(last_err) if last_err else "no PPTX renderer found (install PowerPoint or LibreOffice)")


def render_html(html_path: str, out_dir: str, slides: int, width: int, height: int, tag: str) -> List[str]:
    """Screenshot each slide of the standalone HTML (uses its ?slide=N deep link) in headless Chrome."""
    chrome = next((p for p in CHROME_PATHS if os.path.exists(p)), None)
    if not chrome:
        raise RenderUnavailable("Chrome/Chromium not found")
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for n in range(1, slides + 1):
        out = os.path.join(out_dir, f"html_{tag}_{n:02d}.png")
        url = "file://" + os.path.abspath(html_path) + f"?slide={n}&static=1"
        subprocess.run(
            [chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars",
             f"--window-size={width},{height}", "--virtual-time-budget=6000",
             f"--screenshot={out}", url],
            check=True, capture_output=True, timeout=60,
        )
        paths.append(out)
    return paths


def analyze_png(path: str) -> Dict[str, Any]:
    """Pixel metrics: longest band of rows with no text/graphics in the content area."""
    from PIL import Image

    im = Image.open(path).convert("RGB")
    w, h = im.size
    px = im.load()
    top = int(h * CONTENT_TOP_RATIO)
    step = max(1, w // 400)

    def row_has_ink(y: int) -> bool:
        dark = 0
        for x in range(0, w, step):
            r, g, b = px[x, y]
            if (r + g + b) / 3 < 190 or (max(r, g, b) - min(r, g, b)) > 60:
                dark += 1
                if dark >= 3:
                    return True
        return False

    longest = run = 0
    last_ink = top
    for y in range(top, h):
        if row_has_ink(y):
            run = 0
            last_ink = y
        else:
            run += 1
            longest = max(longest, run)
    return {
        "size": [w, h],
        "blank_band_ratio": round(longest / h, 3),
        "content_bottom_ratio": round(last_ink / h, 3),
    }


def check_pngs(paths: List[str]) -> List[Dict[str, Any]]:
    findings = []
    for i, p in enumerate(paths, 1):
        m = analyze_png(p)
        if m["blank_band_ratio"] > BLANK_BAND_LIMIT:
            findings.append({"slide": i, "code": "BLANK_BAND",
                             "message": f"{m['blank_band_ratio']:.0%} of the slide height is one empty band (limit {BLANK_BAND_LIMIT:.0%})",
                             "image": p})
    return findings


def check_html_layout(html_path: str, slides: int, width: int = 1600, height: int = 900) -> List[Dict[str, Any]]:
    """Open each slide in static mode and read back how far its content extends past the canvas."""
    chrome = next((p for p in CHROME_PATHS if os.path.exists(p)), None)
    if not chrome:
        raise RenderUnavailable("Chrome/Chromium not found")
    findings = []
    for n in range(1, slides + 1):
        url = "file://" + os.path.abspath(html_path) + f"?slide={n}&static=1"
        out = subprocess.run(
            [chrome, "--headless=new", "--disable-gpu", f"--window-size={width},{height}",
             "--virtual-time-budget=6000", "--dump-dom", url],
            check=True, capture_output=True, timeout=60,
        ).stdout.decode("utf-8", "replace")
        m = re.search(r"data-layout=\"([^\"]+)\"", out)
        if not m:
            findings.append({"slide": n, "code": "HTML_NO_REPORT", "message": "page did not report its layout"})
            continue
        info = json.loads(m.group(1).replace("&quot;", '"'))
        if info["bottom"] > 756:
            findings.append({"slide": n, "code": "HTML_OVERFLOW_BOTTOM",
                             "message": f"content extends {info['bottom'] - 754}px below the 754px canvas at {width}x{height}"})
        if info["right"] > 1342:
            findings.append({"slide": n, "code": "HTML_OVERFLOW_RIGHT",
                             "message": f"content extends {info['right'] - 1340}px past the 1340px canvas at {width}x{height}"})
    return findings


def powerpoint_animated_shapes(pptx_path: str, timeout: int = 150) -> Dict[int, int]:
    """Ask PowerPoint how many shapes on each slide it recognises as animated ({slide_number: count}).

    This is the ground truth for motion: an animation tree PowerPoint does not understand reads as 0.
    """
    if not available_renderers()["powerpoint"]:
        raise RenderUnavailable("PowerPoint is not available")
    os.makedirs(STAGING_DIR, exist_ok=True)
    staged = os.path.join(STAGING_DIR, "motion.pptx")
    shutil.copyfile(pptx_path, staged)
    script = f'''
tell application "Microsoft PowerPoint"
  open POSIX file "{staged}"
  delay 2
  set pres to active presentation
  set report to ""
  repeat with n from 1 to (count of slides of pres)
    set sld to slide n of pres
    set animated to 0
    repeat with idx from 1 to (count of shapes of sld)
      try
        if (animate of animation settings of shape idx of sld) is true then set animated to animated + 1
      end try
    end repeat
    set report to report & n & ":" & animated & ";"
  end repeat
  close pres saving no
  return report
end tell
'''
    try:
        out = subprocess.run(["osascript", "-e", script], check=True, capture_output=True, timeout=timeout).stdout.decode()
    except subprocess.SubprocessError as exc:
        raise RenderUnavailable(f"PowerPoint animation probe failed: {exc}") from exc
    finally:
        if os.path.exists(staged):
            os.remove(staged)
    result: Dict[int, int] = {}
    for part in out.strip().split(";"):
        if ":" in part:
            n, c = part.split(":")
            result[int(n)] = int(c)
    return result


def check_motion(pptx_path: str) -> List[Dict[str, Any]]:
    """Compare the animations written into the PPTX with what PowerPoint recognises."""
    from pptx import Presentation

    from core.motion import timing_summary, validate_timing

    prs = Presentation(pptx_path)
    findings: List[Dict[str, Any]] = []
    expected: Dict[int, int] = {}
    for i, slide in enumerate(prs.slides, 1):
        for problem in validate_timing(slide):
            findings.append({"slide": i, "code": "MOTION_INVALID", "message": problem})
        summary = timing_summary(slide)
        if summary["present"]:
            expected[i] = len(summary["spids"])
    if not expected:
        return findings
    recognised = powerpoint_animated_shapes(pptx_path)
    for i, want in expected.items():
        got = recognised.get(i, 0)
        if got < want:
            findings.append({"slide": i, "code": "MOTION_NOT_RECOGNIZED",
                             "message": f"the file animates {want} shape(s) but PowerPoint recognises {got}"})
    return findings
